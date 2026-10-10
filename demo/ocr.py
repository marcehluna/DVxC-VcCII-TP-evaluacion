"""OCR en vivo sobre las máscaras producidas por YOLO26s-seg.

La implementación replica el flujo probado en notebooks/OCR.ipynb sin importar ni
modificar la notebook: recorte orientado, CLAHE, rotaciones, PaddleOCR y matching.
"""

from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from time import perf_counter
from typing import Iterator, Sequence
import os

import cv2
import numpy as np
from PIL import Image
from rapidfuzz import fuzz, process


ROOT = Path(__file__).resolve().parents[1]
PADDLEX_CACHE = ROOT / ".cache" / "paddlex"
PADDLE_DETECTION_MODEL = "PP-OCRv3_mobile_det"
PADDLE_RECOGNITION_MODEL = "latin_PP-OCRv3_mobile_rec"


@dataclass(frozen=True)
class CatalogEntry:
    title: str
    author: str


@dataclass
class OCRReading:
    number: int
    crop: Image.Image
    detection_score: float
    text: str
    confidence: float
    rotation: int
    latency_ms: float
    catalog_title: str | None = None
    catalog_author: str | None = None
    match_score: float = 0.0

    @property
    def matched(self) -> bool:
        return self.catalog_title is not None


def crop_from_polygon(image_bgr: np.ndarray, points: Sequence[Sequence[float]]) -> np.ndarray:
    """Recorta una máscara usando su rectángulo mínimo y endereza el eje largo."""
    polygon = np.asarray(points, dtype=np.float32)
    if polygon.ndim != 2 or polygon.shape[0] < 3 or polygon.shape[1] != 2:
        return np.empty((0, 0, 3), dtype=np.uint8)
    (center_x, center_y), (crop_width, crop_height), angle = cv2.minAreaRect(polygon)
    if crop_width < crop_height:
        crop_width, crop_height = crop_height, crop_width
        angle += 90

    matrix = cv2.getRotationMatrix2D((center_x, center_y), angle, 1.0)
    height, width = image_bgr.shape[:2]
    rotated = cv2.warpAffine(
        image_bgr,
        matrix,
        (width, height),
        flags=cv2.INTER_CUBIC,
        borderMode=cv2.BORDER_REPLICATE,
    )
    x1 = max(0, int(center_x - crop_width / 2))
    y1 = max(0, int(center_y - crop_height / 2))
    x2 = min(width, int(center_x + crop_width / 2))
    y2 = min(height, int(center_y + crop_height / 2))
    return rotated[y1:y2, x1:x2]


def preprocess_crop(crop_bgr: np.ndarray, min_width: int = 80) -> np.ndarray:
    """Aumenta recortes angostos y mejora el contraste antes de leerlos."""
    if crop_bgr.size == 0:
        return crop_bgr
    height, width = crop_bgr.shape[:2]
    if 0 < width < min_width:
        scale = min_width / width
        crop_bgr = cv2.resize(
            crop_bgr,
            (round(width * scale), round(height * scale)),
            interpolation=cv2.INTER_CUBIC,
        )
    lab = cv2.cvtColor(crop_bgr, cv2.COLOR_BGR2LAB)
    lightness, channel_a, channel_b = cv2.split(lab)
    lightness = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(lightness)
    return cv2.cvtColor(cv2.merge([lightness, channel_a, channel_b]), cv2.COLOR_LAB2BGR)


@lru_cache(maxsize=1)
def paddle_reader():
    """Carga una sola vez los modelos de detección y reconocimiento de PaddleOCR."""
    PADDLEX_CACHE.mkdir(parents=True, exist_ok=True)
    os.environ.setdefault("PADDLE_PDX_CACHE_HOME", str(PADDLEX_CACHE))
    os.environ.setdefault("PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK", "True")
    from paddleocr import PaddleOCR

    # Son los modelos que PaddleOCR seleccionó para lang="es" en el experimento.
    # La orientación se resuelve explícitamente abajo, igual que en la notebook.
    return PaddleOCR(
        text_detection_model_name=PADDLE_DETECTION_MODEL,
        text_recognition_model_name=PADDLE_RECOGNITION_MODEL,
        device="cpu",
        enable_mkldnn=False,
        use_doc_orientation_classify=False,
        use_doc_unwarping=False,
        use_textline_orientation=False,
    )


def _result_value(result, key: str):
    try:
        return result[key]
    except (KeyError, TypeError, IndexError):
        return getattr(result, key, None)


def read_paddle(crop_bgr: np.ndarray) -> tuple[str, float]:
    results = list(
        paddle_reader().predict(
            crop_bgr,
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=False,
        )
    )
    if not results:
        return "", 0.0
    texts = list(_result_value(results[0], "rec_texts") or [])
    scores = list(_result_value(results[0], "rec_scores") or [])
    if not texts:
        return "", 0.0
    confidence = float(np.mean(scores)) if scores else 0.0
    return " ".join(str(text).strip() for text in texts if str(text).strip()).strip(), confidence


def read_best_rotation(
    crop_bgr: np.ndarray,
    rotations: Sequence[int] = (0, 90, 270),
) -> tuple[str, float, int, float, np.ndarray]:
    """Prueba las orientaciones del experimento y conserva la lectura más confiable."""
    start = perf_counter()
    best_text, best_confidence, best_rotation = "", -1.0, 0
    best_image = crop_bgr
    for rotation in rotations:
        if rotation == 0:
            rotated = crop_bgr
        elif rotation == 90:
            rotated = cv2.rotate(crop_bgr, cv2.ROTATE_90_CLOCKWISE)
        elif rotation == 270:
            rotated = cv2.rotate(crop_bgr, cv2.ROTATE_90_COUNTERCLOCKWISE)
        else:
            raise ValueError(f"Rotación OCR no soportada: {rotation}")
        text, confidence = read_paddle(rotated)
        if confidence > best_confidence:
            best_text, best_confidence, best_rotation = text, confidence, rotation
            best_image = rotated
    return best_text, max(0.0, best_confidence), best_rotation, (perf_counter() - start) * 1000, best_image


@lru_cache(maxsize=4)
def _load_catalog(path: str, modified_ns: int) -> tuple[CatalogEntry, ...]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip().lower() for value in next(rows)]
        title_index = headers.index("titulo")
        author_index = headers.index("autor") if "autor" in headers else None
        entries = []
        for row in rows:
            title = str(row[title_index] or "").strip()
            if not title:
                continue
            author = str(row[author_index] or "").strip() if author_index is not None else ""
            entries.append(CatalogEntry(title, author))
        return tuple(entries)
    finally:
        workbook.close()


def catalog_entries(path: Path = ROOT / "catalogo" / "catalogo.xlsx") -> tuple[CatalogEntry, ...]:
    if not path.is_file():
        return ()
    return _load_catalog(str(path.resolve()), path.stat().st_mtime_ns)


def match_catalog(
    text: str,
    entries: Sequence[CatalogEntry],
    threshold: float = 85.0,
) -> tuple[str | None, str | None, float]:
    if not text.strip() or not entries:
        return None, None, 0.0
    titles = [entry.title.upper() for entry in entries]
    match = process.extractOne(text.upper(), titles, scorer=fuzz.token_set_ratio)
    if match is None:
        return None, None, 0.0
    _, score, index = match
    entry = entries[index]
    if score < threshold:
        return None, None, float(score)
    return entry.title, entry.author, float(score)


def _selected_polygons(
    polygons: Sequence[Sequence[Sequence[float]]],
    scores: Sequence[float],
    limit: int,
) -> list[tuple[Sequence[Sequence[float]], float]]:
    ranked = sorted(zip(polygons, scores), key=lambda item: item[1], reverse=True)[:limit]
    # Orden de lectura aproximado: de izquierda a derecha sobre el estante.
    return sorted(ranked, key=lambda item: float(np.asarray(item[0])[:, 0].mean()))


def iter_readings(
    image: Image.Image,
    polygons: Sequence[Sequence[Sequence[float]]],
    scores: Sequence[float],
    limit: int = 20,
    match_threshold: float = 85.0,
) -> Iterator[OCRReading]:
    """Procesa los lomos con más confianza y entrega resultados incrementalmente."""
    rgb = np.asarray(image.convert("RGB"))
    image_bgr = cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR)
    entries = catalog_entries()
    for number, (polygon, detection_score) in enumerate(
        _selected_polygons(polygons, scores, max(1, int(limit))), start=1
    ):
        crop = preprocess_crop(crop_from_polygon(image_bgr, polygon))
        if crop.size == 0:
            continue
        text, confidence, rotation, latency_ms, oriented = read_best_rotation(crop)
        title, author, match_score = match_catalog(text, entries, match_threshold)
        yield OCRReading(
            number=number,
            crop=Image.fromarray(cv2.cvtColor(oriented, cv2.COLOR_BGR2RGB)),
            detection_score=float(detection_score),
            text=text,
            confidence=confidence,
            rotation=rotation,
            latency_ms=latency_ms,
            catalog_title=title,
            catalog_author=author,
            match_score=match_score,
        )
